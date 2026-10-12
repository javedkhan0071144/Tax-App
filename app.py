
import streamlit as st
import pandas as pd
from pathlib import Path
from openpyxl import Workbook, load_workbook
from datetime import datetime
import shutil

BASE = Path(__file__).parent
DATA_FILE = BASE / "professional_tax_data.xlsx"
BACKUP_DIR = BASE / "backups"

COLUMNS = [
    "Serial No", "Business Name", "Category", "Sub Category",
    "Financial From", "Financial To", "Current Year Amount",
    "Arrears Amount", "Total Amount", "Last Date"
]

st.set_page_config(page_title="Professional Tax System", layout="wide")


def setup_excel():
    if not DATA_FILE.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "Taxpayers"
        ws.append(COLUMNS)
        wb.save(DATA_FILE)
        wb.close()


def read_data():
    setup_excel()
    return pd.read_excel(DATA_FILE).fillna("")


def save_record(record, serial=None):
    setup_excel()
    wb = load_workbook(DATA_FILE)
    ws = wb["Taxpayers"]

    if serial is None:
        serials = [
            int(ws.cell(r, 1).value)
            for r in range(2, ws.max_row + 1)
            if str(ws.cell(r, 1).value).isdigit()
        ]
        serial = max(serials, default=0) + 1
        record["Serial No"] = serial
        ws.append([record.get(c, "") for c in COLUMNS])
    else:
        for r in range(2, ws.max_row + 1):
            if str(ws.cell(r, 1).value) == str(serial):
                for i, c in enumerate(COLUMNS, 1):
                    ws.cell(r, i, record.get(c, ""))
                break

    wb.save(DATA_FILE)
    wb.close()


def backup_data():
    if DATA_FILE.exists():
        BACKUP_DIR.mkdir(exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        shutil.copy2(DATA_FILE, BACKUP_DIR / f"backup_{stamp}.xlsx")


setup_excel()
df = read_data()

st.title("Professional Tax Management System")
st.sidebar.title("Main Menu")
page = st.sidebar.radio(
    "Select page",
    ["Dashboard", "Add Taxpayer", "Search / Edit", "Reports", "Backup"]
)

if page == "Dashboard":
    st.header("Dashboard")
    a, b, c = st.columns(3)
    a.metric("Total Taxpayers", len(df))
    current = pd.to_numeric(df["Current Year Amount"], errors="coerce").fillna(0).sum()
    total = pd.to_numeric(df["Total Amount"], errors="coerce").fillna(0).sum()
    b.metric("Current Tax Assessed", f"Rs. {current:,.0f}")
    c.metric("Total Tax and Arrears", f"Rs. {total:,.0f}")
    st.dataframe(df, use_container_width=True)

elif page == "Add Taxpayer":
    st.header("Add New Taxpayer")
    with st.form("add"):
        business = st.text_input("Business Name")
        category = st.text_input("Category")
        subcategory = st.text_input("Sub Category")
        col1, col2 = st.columns(2)
        year_from = col1.text_input("Financial From", "2026")
        year_to = col2.text_input("Financial To", "2027")
        col3, col4 = st.columns(2)
        current = col3.number_input("Current Year Amount", min_value=0.0)
        arrears = col4.number_input("Arrears Amount", min_value=0.0)
        last_date = st.date_input("Last Date")
        submitted = st.form_submit_button("Save Taxpayer")

    if submitted:
        if not business.strip():
            st.error("Enter the Business Name.")
        else:
            save_record({
                "Business Name": business.strip(),
                "Category": category,
                "Sub Category": subcategory,
                "Financial From": year_from,
                "Financial To": year_to,
                "Current Year Amount": current,
                "Arrears Amount": arrears,
                "Total Amount": current + arrears,
                "Last Date": last_date.isoformat(),
            })
            backup_data()
            st.success("Taxpayer saved to Excel. Backup created.")
            st.rerun()

elif page == "Search / Edit":
    st.header("Search / Edit Taxpayer")
    if df.empty:
        st.info("No records yet. Add a taxpayer first.")
    else:
        search = st.text_input("Search by business name or serial number")
        filtered = df[df.astype(str).apply(
            lambda col: col.str.contains(search, case=False, na=False)
        ).any(axis=1)] if search else df
        st.dataframe(filtered, use_container_width=True)

        if not filtered.empty:
            selected = st.selectbox("Select Serial No", filtered["Serial No"].tolist())
            old = df[df["Serial No"].astype(str) == str(selected)].iloc[0]
            with st.form("edit"):
                business = st.text_input("Business Name", str(old["Business Name"]))
                category = st.text_input("Category", str(old["Category"]))
                subcategory = st.text_input("Sub Category", str(old["Sub Category"]))
                year_from = st.text_input("Financial From", str(old["Financial From"]))
                year_to = st.text_input("Financial To", str(old["Financial To"]))
                current = st.number_input("Current Year Amount", min_value=0.0,
                                           value=float(old["Current Year Amount"] or 0))
                arrears = st.number_input("Arrears Amount", min_value=0.0,
                                          value=float(old["Arrears Amount"] or 0))
                last_date = st.text_input("Last Date", str(old["Last Date"]))
                updated = st.form_submit_button("Update Record")

            if updated:
                save_record({
                    "Business Name": business,
                    "Category": category,
                    "Sub Category": subcategory,
                    "Financial From": year_from,
                    "Financial To": year_to,
                    "Current Year Amount": current,
                    "Arrears Amount": arrears,
                    "Total Amount": current + arrears,
                    "Last Date": last_date,
                }, serial=selected)
                backup_data()
                st.success("Record updated and backup created.")
                st.rerun()

elif page == "Reports":
    st.header("Reports")
    st.dataframe(df, use_container_width=True)
    st.download_button(
        "Download Excel Data",
        DATA_FILE.read_bytes(),
        file_name="professional_tax_data.xlsx"
    )
    st.download_button(
        "Download CSV Report",
        df.to_csv(index=False).encode("utf-8-sig"),
        file_name="professional_tax_report.csv"
    )

elif page == "Backup":
    st.header("Backup Management")
    if st.button("Create Manual Backup"):
        backup_data()
        st.success("Backup created.")
    if BACKUP_DIR.exists():
        for f in sorted(BACKUP_DIR.glob("*.xlsx"), reverse=True):
            st.write(f.name)
    else:
        st.info("No backups yet.")
